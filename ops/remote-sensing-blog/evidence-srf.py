"""Read actual ESA SRF workbook, including worksheet names and sample spacing."""
from pathlib import Path
import urllib.request
import openpyxl
import json
import hashlib
from PIL import Image, ImageDraw, ImageFont
ROOT=Path(__file__).resolve().parent
url='https://sentiwiki.copernicus.eu/__attachments/a_ece6183b6698587e7ecd9804974bece3d82e39e151aa5aeca9d8fd95e1f1558c/COPE-GSEG-EOPG-TN-15-0007%20-%20Sentinel-2%20Spectral%20Response%20Functions%202024%20-%204.0.xlsx'
path=ROOT/'cache/s2-srf-v4.xlsx'
if not path.exists():
 with urllib.request.urlopen(url,timeout=55) as r:
  data=r.read();assert data.startswith(b'PK');path.write_bytes(data)
wb=openpyxl.load_workbook(path,data_only=True,read_only=False)
for ws in wb:
 print('\nSHEET',ws.title,ws.max_row,ws.max_column)
 for row in list(ws.values)[:9]:print(row)
 for i,im in enumerate(ws._images):
  filename=ROOT/'cache'/('srf-'+ws.title.replace(' ','-')+f'-{i}.'+im.format)
  filename.write_bytes(im._data())
  print('ORIGINAL EMBEDDED IMAGE',filename)

out=ROOT.parents[1]/'src/content/blog/flaash-atmospheric-correction-concepts/evidence-32.png'
parts=[]
sources=[]
for title,sheet in [('等效波长：对响应函数进行波长加权','Equivalent Wavelengths'),('带宽与带宽中点：由两侧最陡斜率的位置定义','Bandwidth and mid-wavelength')]:
 im=wb[sheet]._images[0]
 source=ROOT/'cache'/('srf-'+sheet.replace(' ','-')+'-0.'+im.format)
 original=Image.open(source).convert('RGB')
 draw=ImageDraw.Draw(original)
 if sheet=='Equivalent Wavelengths': draw.rectangle((80,44,330,143),outline='#cf202f',width=2)
 else: draw.rectangle((67,353,419,411),outline='#cf202f',width=2)
 original=original.resize((original.width*2,original.height*2),Image.Resampling.LANCZOS)
 parts.append((title,sheet,original))
 sources.append({'worksheet':sheet,'anchor_cell':[im.anchor._from.row+1,im.anchor._from.col+1],'original_image_sha256':hashlib.sha256(source.read_bytes()).hexdigest()})
canvas=Image.new('RGB',(1320,150+sum(95+im.height for _,_,im in parts)+35),'#f6f7fa')
d=ImageDraw.Draw(canvas)
font=lambda s:ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',s)
d.text((45,25),'原文证据 32 · 官方光谱响应表里的两个不同定义',font=font(29),fill='#253248')
d.text((45,76),'原 XLSX 内嵌附图直接提取；仅加中文说明、红框和等比例放大。',font=font(22),fill='#697386')
y=139
for title,sheet,im in parts:
 d.text((45,y),title,font=font(23),fill='#253248');y+=37
 d.text((45,y),'MSI Spectral Response Functions 4.0 · 工作表 '+sheet,font=font(20),fill='#697386');y+=32
 canvas.paste(im,(56,y));y+=im.height+26
canvas.save(out,optimize=True)
ws=wb['Spectral Responses (S2A)']
wavelengths=[row[0] for row in list(ws.values)[1:] if isinstance(row[0],(int,float))]
assert all(b-a==1 for a,b in zip(wavelengths,wavelengths[1:]))
report={'source_url':url,'file_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'version':wb['Overview']['F6'].value,'worksheet_rows':'Spectral Responses (S2A)!A2:A2302','wavelength_step_nm':1,'wavelength_extent_nm':[min(wavelengths),max(wavelengths)],'original_embedded_figures':sources,'rendered_file':out.name}
(ROOT/'evidence-srf.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print('SRF evidence saved',report)
